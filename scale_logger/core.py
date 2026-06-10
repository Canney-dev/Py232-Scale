import math
import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet


PROJECT_DIR = Path(__file__).resolve().parent.parent
LOG_FOLDER = PROJECT_DIR / "Logs"
TEMPLATE_FILENAME = PROJECT_DIR / "weight_qc_checklist_template.xlsx"

DEFAULT_PORT = "COM3"
DEFAULT_BAUDRATE = 9600
FORM_START_ROW = 6
FORM_END_ROW = 34
WEIGHTS_PER_ITEM = 3
LOW_SUMMARY_CELL = "E37"
HIGH_SUMMARY_CELL = "E38"

SCALE_PATTERN = re.compile(
    r"([+-]?(?:\d+(?:\.\d+)?|\.\d+))\s*([A-Z]{0,4})(?:\s*([LGH]))?",
    re.IGNORECASE,
)
NET_WEIGHT_PATTERN = re.compile(
    r"\bNET\s*:\s*([+-]?(?:\d+(?:\.\d+)?|\.\d+))\s*([A-Z]{0,4})",
    re.IGNORECASE,
)
IGNORED_TICKET_LABELS = (
    "DATE:",
    "TIME:",
    "GROSS:",
    "TARE:",
    "MERCHANDISE:",
    "PIECE WEIGHT:",
    "COUNT:",
)


def clean_scale_line(line):
    return "".join(char for char in line if char.isprintable()).strip()


@dataclass
class Item:
    name: str
    description: str


def caps(value):
    return str(value).upper()


def title_text(value):
    return str(value).strip().title()


def parse_scale_line(line):
    line = clean_scale_line(line)
    net_match = NET_WEIGHT_PATTERN.search(line)

    if not net_match and line.upper().startswith(IGNORED_TICKET_LABELS):
        return None

    match = net_match or SCALE_PATTERN.search(line)

    if not match:
        return None

    weight = float(match.group(1))
    unit = (match.group(2) or "").upper()
    status_code = "" if net_match else (match.group(3) or "").upper()
    status_map = {
        "L": "Low",
        "G": "Good",
        "H": "High",
    }

    return {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "weight": weight,
        "unit": unit,
        "status_code": status_code,
        "status": status_map.get(status_code, "Zero" if is_zero(weight) else "Unknown"),
        "raw": line,
    }


def is_zero(weight):
    return weight == 0


def get_next_log_filename():
    os.makedirs(LOG_FOLDER, exist_ok=True)

    existing = [
        name for name in os.listdir(LOG_FOLDER)
        if name.startswith("Weight_QC_Form ") and name.endswith(".xlsx")
    ]
    numbers = []

    for name in existing:
        try:
            number = int(name.replace("Weight_QC_Form ", "").replace(".xlsx", ""))
            numbers.append(number)
        except ValueError:
            pass

    next_number = max(numbers, default=0) + 1
    return LOG_FOLDER / f"Weight_QC_Form {next_number}.xlsx"


def apply_qc_form_layout(sheet):
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E7E7E7")

    sheet.title = "QC Form"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_area = "A1:E40"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.30
    sheet.page_margins.bottom = 0.30
    sheet.page_margins.header = 0.15
    sheet.page_margins.footer = 0.15
    sheet.freeze_panes = "A6"

    widths = {
        "A": 12,
        "B": 34,
        "C": 11,
        "D": 11,
        "E": 11,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in range(1, 41):
        sheet.row_dimensions[row].height = 15

    for row in range(FORM_START_ROW, FORM_END_ROW + 1):
        sheet.row_dimensions[row].height = 18

    sheet.row_dimensions[4].height = 21
    sheet.row_dimensions[40].height = 20

    sheet["A2"] = "Project #"
    sheet["A2"].font = Font(bold=True)
    sheet["B2"].border = Border(bottom=medium)
    sheet.merge_cells("B2:D2")

    sheet.merge_cells("A4:E4")
    sheet["A4"] = "Weight QC Checklist Form"
    sheet["A4"].font = Font(bold=True, size=14)
    sheet["A4"].alignment = Alignment(horizontal="center")

    headers = ["Item", "Description", "Weight 1", "Weight 2", "Weight 3"]
    for cell, header in zip(sheet[5], headers):
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = table_border

    for row in range(FORM_START_ROW, FORM_END_ROW + 1):
        for column in range(1, 6):
            cell = sheet.cell(row=row, column=column)
            cell.border = table_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet.merge_cells("C37:D37")
    sheet["C37"] = "Acceptable Tolerance"
    sheet["C37"].font = Font(bold=True)
    sheet["C37"].alignment = Alignment(horizontal="center")

    for cell_name, label in ((LOW_SUMMARY_CELL, "Low"), (HIGH_SUMMARY_CELL, "High")):
        cell = sheet[cell_name]
        cell.value = label
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=medium)

    sheet["B40"] = "Line Lead Sign Off"
    sheet["B40"].font = Font(bold=True)
    sheet["C40"].border = Border(bottom=medium)
    sheet.merge_cells("C40:E40")


def adjust_form_row_height(sheet, row):
    max_lines = 1

    for column in range(1, 6):
        cell = sheet.cell(row=row, column=column)
        value = "" if cell.value is None else str(cell.value)
        column_letter = get_column_letter(column)
        column_width = sheet.column_dimensions[column_letter].width or 10
        usable_width = max(1, int(column_width) - 2)
        wrapped_lines = 0

        for line in value.splitlines() or [""]:
            wrapped_lines += max(1, math.ceil(len(line) / usable_width))

        max_lines = max(max_lines, wrapped_lines)

    sheet.row_dimensions[row].height = max(18, min(72, max_lines * 15))


def update_summary(sheet, cell_name, label, entries):
    summary_cell = sheet[cell_name]
    assert isinstance(summary_cell, Cell)

    if entries:
        summary_cell.value = f"{label}\n" + "\n".join(entries)
    else:
        summary_cell.value = label

    summary_cell.alignment = Alignment(vertical="top", wrap_text=True)
    row = summary_cell.row
    sheet.row_dimensions[row].height = max(20, min(90, (len(entries) + 1) * 15))


def create_qc_workbook(project_number="", items=None):
    workbook = Workbook()
    form_sheet = workbook.active
    assert isinstance(form_sheet, Worksheet)

    apply_qc_form_layout(form_sheet)
    form_sheet["B2"] = title_text(project_number)

    if items:
        for index, item in enumerate(items, start=FORM_START_ROW):
            form_sheet.cell(row=index, column=1).value = title_text(item.name)
            form_sheet.cell(row=index, column=2).value = title_text(item.description)
            adjust_form_row_height(form_sheet, index)

    raw_sheet = workbook.create_sheet("Raw Data")
    raw_sheet.append([
        "Timestamp",
        "Item",
        "Description",
        "Weight Number",
        "Weight",
        "Unit",
        "Status Code",
        "Scale Status",
        "QC Status",
        "Raw",
    ])
    for column in range(1, 11):
        cell = raw_sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7E7E7")
        raw_sheet.column_dimensions[get_column_letter(column)].width = 18

    raw_sheet.column_dimensions["C"].width = 30
    raw_sheet.column_dimensions["J"].width = 30

    return workbook


def rank_weight_statuses(readings):
    weights = [reading["weight"] for reading in readings]
    low_index = min(range(len(weights)), key=weights.__getitem__)
    high_index = max(range(len(weights)), key=weights.__getitem__)

    if low_index == high_index and len(weights) > 1:
        high_index = len(weights) - 1

    statuses = [""] * len(weights)
    statuses[low_index] = "Low"
    statuses[high_index] = "High"

    return statuses


class QcWorkbookSession:
    def __init__(self, project_number, items):
        self.project_number = title_text(project_number)
        self.items = [Item(title_text(item.name), title_text(item.description)) for item in items]
        self.filename = get_next_log_filename()
        self.workbook = create_qc_workbook(self.project_number, self.items)
        self.form_sheet = self.workbook["QC Form"]
        self.raw_sheet = self.workbook["Raw Data"]
        self.low_entries = []
        self.high_entries = []
        self.item_results = {}
        self.save()

    def save(self):
        self.workbook.save(self.filename)

    @classmethod
    def from_file(cls, filename):
        session = cls.__new__(cls)
        session.filename = Path(filename)
        session.workbook = load_workbook(session.filename)
        session.form_sheet = (
            session.workbook["QC Form"]
            if "QC Form" in session.workbook.sheetnames
            else session.workbook["QC FORM"]
        )
        session.raw_sheet = (
            session.workbook["Raw Data"]
            if "Raw Data" in session.workbook.sheetnames
            else session.workbook["RAW DATA"]
        )
        session.project_number = title_text(session.form_sheet["B2"].value or "")
        session.items = []
        session.low_entries = []
        session.high_entries = []
        session.item_results = {}

        for row in range(FORM_START_ROW, FORM_END_ROW + 1):
            name = session.form_sheet.cell(row=row, column=1).value
            description = session.form_sheet.cell(row=row, column=2).value
            if name is None and description is None:
                continue

            session.items.append(Item(title_text(name or ""), title_text(description or "")))

        item_lookup = {}
        for index, item in enumerate(session.items):
            item_lookup.setdefault(item.name, index)

        for raw_row in range(2, session.raw_sheet.max_row + 1):
            item_name = title_text(session.raw_sheet.cell(row=raw_row, column=2).value or "")
            if item_name not in item_lookup:
                continue

            weight_number = session.raw_sheet.cell(row=raw_row, column=4).value
            weight = session.raw_sheet.cell(row=raw_row, column=5).value
            if weight_number is None or weight is None:
                continue

            item_index = item_lookup[item_name]
            result = session.item_results.setdefault(
                item_index,
                {"readings": [None] * WEIGHTS_PER_ITEM, "raw_rows": [None] * WEIGHTS_PER_ITEM},
            )
            reading_index = int(weight_number) - 1
            if reading_index < 0 or reading_index >= WEIGHTS_PER_ITEM:
                continue

            result["readings"][reading_index] = {
                "timestamp": caps(session.raw_sheet.cell(row=raw_row, column=1).value or ""),
                "weight": float(weight),
                "unit": caps(session.raw_sheet.cell(row=raw_row, column=6).value or ""),
                "status_code": caps(session.raw_sheet.cell(row=raw_row, column=7).value or ""),
                "status": title_text(session.raw_sheet.cell(row=raw_row, column=8).value or ""),
                "raw": str(session.raw_sheet.cell(row=raw_row, column=10).value or ""),
            }
            result["raw_rows"][reading_index] = raw_row

        incomplete_items = []
        for item_index, result in session.item_results.items():
            readings = result["readings"]
            raw_rows = result["raw_rows"]
            if any(reading is None or raw_row is None for reading, raw_row in zip(readings, raw_rows)):
                incomplete_items.append(item_index)

        for item_index in incomplete_items:
            del session.item_results[item_index]

        session.rebuild_summaries()
        return session

    def update_item(self, item_index, name, description):
        item = self.items[item_index]
        item.name = title_text(name)
        item.description = title_text(description)

        form_row = FORM_START_ROW + item_index
        self.form_sheet.cell(row=form_row, column=1).value = item.name
        self.form_sheet.cell(row=form_row, column=2).value = item.description
        adjust_form_row_height(self.form_sheet, form_row)

        if item_index in self.item_results:
            for raw_row in self.item_results[item_index]["raw_rows"]:
                self.raw_sheet.cell(row=raw_row, column=2).value = item.name
                self.raw_sheet.cell(row=raw_row, column=3).value = item.description

            self.rebuild_summaries()

        self.save()

    def record_reading(self, item_index, weight_number, reading):
        item = self.items[item_index]
        form_row = FORM_START_ROW + item_index
        form_cell = self.form_sheet.cell(row=form_row, column=weight_number + 2)
        assert isinstance(form_cell, Cell)
        form_cell.value = self.format_weight(reading)
        form_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        adjust_form_row_height(self.form_sheet, form_row)

        self.raw_sheet.append([
            caps(reading["timestamp"]),
            item.name,
            item.description,
            weight_number,
            reading["weight"],
            caps(reading["unit"]),
            caps(reading["status_code"]),
            title_text(reading["status"]),
            "",
            str(reading["raw"]),
        ])
        raw_row = self.raw_sheet.max_row
        self.save()
        return raw_row

    def finalize_item(self, item_index, readings, raw_rows):
        form_row = FORM_START_ROW + item_index
        statuses = rank_weight_statuses(readings)
        self.item_results[item_index] = {
            "readings": readings,
            "raw_rows": raw_rows,
        }

        for column in range(3, 6):
            cell = self.form_sheet.cell(row=form_row, column=column)
            assert isinstance(cell, Cell)
            cell.value = None

        for weight_number, (reading, status, raw_row) in enumerate(
            zip(readings, statuses, raw_rows),
            start=1,
        ):
            form_cell = self.form_sheet.cell(row=form_row, column=weight_number + 2)
            assert isinstance(form_cell, Cell)
            value = self.format_weight(reading)
            if status:
                value += f"\n{status}"
            form_cell.value = value
            form_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            adjust_form_row_height(self.form_sheet, form_row)

            raw_status_cell = self.raw_sheet.cell(row=raw_row, column=9)
            assert isinstance(raw_status_cell, Cell)
            raw_status_cell.value = status

        self.rebuild_summaries()
        self.save()
        return statuses

    def update_reading(self, item_index, weight_number, new_weight, unit):
        result = self.item_results[item_index]
        reading = result["readings"][weight_number - 1]
        raw_row = result["raw_rows"][weight_number - 1]
        reading["weight"] = new_weight
        reading["unit"] = caps(unit)
        reading["raw"] = f"Manual Edit {new_weight:.3f} {caps(unit)}".strip()
        reading["status"] = "Manual Edit"
        reading["status_code"] = ""

        self.raw_sheet.cell(row=raw_row, column=5).value = new_weight
        self.raw_sheet.cell(row=raw_row, column=6).value = caps(unit)
        self.raw_sheet.cell(row=raw_row, column=8).value = "Manual Edit"
        self.raw_sheet.cell(row=raw_row, column=10).value = str(reading["raw"])

        return self.finalize_item(item_index, result["readings"], result["raw_rows"])

    def rebuild_summaries(self):
        low_entries = []
        high_entries = []

        for item_index in sorted(self.item_results):
            item = self.items[item_index]
            readings = self.item_results[item_index]["readings"]
            statuses = rank_weight_statuses(readings)

            for weight_number, (reading, status) in enumerate(
                zip(readings, statuses),
                start=1,
            ):
                entry = f"{item.name} W{weight_number}: {self.format_weight(reading)}"
                if status == "Low":
                    low_entries.append(entry)
                elif status == "High":
                    high_entries.append(entry)

        self.low_entries = low_entries
        self.high_entries = high_entries
        update_summary(self.form_sheet, LOW_SUMMARY_CELL, "Low", self.low_entries)
        update_summary(self.form_sheet, HIGH_SUMMARY_CELL, "High", self.high_entries)

    @staticmethod
    def format_weight(reading):
        unit = caps(reading.get("unit", ""))
        suffix = f" {unit}" if unit else ""
        return f"{reading['weight']:.3f}{suffix}"


def regenerate_template():
    workbook = create_qc_workbook()
    workbook.save(TEMPLATE_FILENAME)
    return TEMPLATE_FILENAME
